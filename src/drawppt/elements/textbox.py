"""Textbox element."""

from dataclasses import dataclass
from typing import Optional

from pptx.util import Pt, Emu
from pptx.enum.text import PP_ALIGN

from .base import Element, px_to_emu


@dataclass
class Textbox(Element):
    """Text box element."""

    text: str = ""
    font_size: int = 18
    font_family: str = "微软雅黑"
    color: str = "#000000"
    bold: bool = False
    italic: bool = False
    align: str = "left"
    line_spacing: float = 1.0
    paragraph_spacing: int = 0

    def __post_init__(self):
        self.element_type = "textbox"

    def apply_to_slide(self, pptx_slide) -> None:
        """Apply textbox to slide."""
        left = px_to_emu(self.x)
        top = px_to_emu(self.y)
        width = px_to_emu(self.w)
        height = px_to_emu(self.h)

        txBox = pptx_slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True

        # Split text into paragraphs
        lines = self.text.split("\n")

        for i, line in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()

            p.text = line
            p.alignment = self._get_alignment()

            # Line spacing
            p.line_spacing = Pt(self.font_size * self.line_spacing)

            # Paragraph spacing (except first paragraph)
            if i > 0 and self.paragraph_spacing > 0:
                p.space_before = Pt(self.paragraph_spacing)

            # Font properties
            for run in p.runs:
                run.font.size = Pt(self.font_size)
                run.font.name = self.font_family
                run.font.bold = self.bold
                run.font.italic = self.italic
                run.font.color.rgb = self._hex_to_rgb(self.color)

    def _get_alignment(self) -> int:
        """Get text alignment enum."""
        align_map = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }
        return align_map.get(self.align, PP_ALIGN.LEFT)

    def _hex_to_rgb(self, hex_color: str):
        """Convert hex color to RGBColor."""
        from pptx.dml.color import RGBColor

        hex_color = hex_color.lstrip("#")
        r = int(hex_color[0:2], 16)
        g = int(hex_color[2:4], 16)
        b = int(hex_color[4:6], 16)
        return RGBColor(r, g, b)

    def to_dict(self) -> dict:
        """Convert to dictionary."""
        base = super().to_dict()
        base.update({
            "text": self.text,
            "font_size": self.font_size,
            "font_family": self.font_family,
            "color": self.color,
            "bold": self.bold,
            "italic": self.italic,
            "align": self.align,
            "line_spacing": self.line_spacing,
            "paragraph_spacing": self.paragraph_spacing,
        })
        return base
